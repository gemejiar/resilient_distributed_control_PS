"""Run E3: replace only SG32 with one REGF1 GFM and validate C0a."""

from __future__ import annotations

import csv
import json
import platform
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import andes
import matplotlib
import numpy as np
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

from run_e0_ieee39 import (  # noqa: E402
    ARTIFACTS_DIR, CRITICAL_BUSES, E0_CASE, FIGURES_DIR, MATRIX_FILE,
    NOMINAL_FREQUENCY_HZ, PROJECT_ROOT, extract_time_series, git_commit, sha256,
)

VARIANT_CASE = E0_CASE.parent / "ieee39_C0a_9SG_1GFM.xlsx"
RAW_FILE = ARTIFACTS_DIR / "IEEE39_E3_C0a_timeseries.npz"
SUMMARY_FILE = ARTIFACTS_DIR / "IEEE39_E3_C0a_summary.json"
FREQUENCY_FIGURE = FIGURES_DIR / "IEEE39_E3_C0a_frequency.png"
VOLTAGE_FIGURE = FIGURES_DIR / "IEEE39_E3_C0a_voltage.png"
POWER_FIGURE = FIGURES_DIR / "IEEE39_E3_C0a_gfm_power.png"

EVENT_TIME_S = 5.0
FINAL_TIME_S = 30.0
FLAT_FINAL_TIME_S = 10.0
TIME_STEP_S = 0.00390625
REFINED_TIME_STEP_S = TIME_STEP_S / 2
STEP_FRACTION = 0.01
GFM_BUS = 32
GFM_STATIC_GEN = 3
GFM_SN_MVA = 843.7
P_EVENT_BUS = 4
Q_EVENT_BUS = 21

REMOVED_DYNAMIC_ROWS = {
    "GENROU": "GENROU_3", "TGOV1N": "TGOV1_3",
    "IEEEX1": "IEEEX1_3", "IEEEST": "IEEEST_3",
}


def build_variant() -> None:
    """Create C0a from the immutable E0 case, removing only SG32 dynamics."""
    if VARIANT_CASE.exists():
        return
    shutil.copy2(E0_CASE, VARIANT_CASE)
    workbook = load_workbook(VARIANT_CASE)
    for sheet_name, idx in REMOVED_DYNAMIC_ROWS.items():
        sheet = workbook[sheet_name]
        idx_col = [cell.value for cell in sheet[1]].index("idx") + 1
        matches = [row for row in range(2, sheet.max_row + 1)
                   if str(sheet.cell(row, idx_col).value) == idx]
        if len(matches) != 1:
            raise RuntimeError(f"Expected exactly one {idx} row; found {matches}.")
        sheet.delete_rows(matches[0], 1)
    sheet = workbook.create_sheet("REGF1")
    # All omitted REGF1 parameters deliberately retain ANDES 2.0.0 defaults.
    sheet.append(["uid", "idx", "u", "name", "bus", "gen", "Sn", "fn"])
    sheet.append([0, "REGF1_32", 1, "GFM32", GFM_BUS, GFM_STATIC_GEN,
                  GFM_SN_MVA, NOMINAL_FREQUENCY_HZ])
    workbook.save(VARIANT_CASE)


def configure_system(scenario: str, time_step_s: float) -> tuple[Any, dict[str, float]]:
    system = andes.load(str(VARIANT_CASE), setup=False, default_config=True, no_output=True)
    system.PQ.config.p2p, system.PQ.config.p2i, system.PQ.config.p2z = 1.0, 0.0, 0.0
    system.PQ.config.q2q, system.PQ.config.q2i, system.PQ.config.q2z = 1.0, 0.0, 0.0
    event = {"p_pu": 0.0, "q_pu": 0.0, "bus": 0.0}
    if scenario != "flat":
        channel, sign = scenario
        source = "Ppf" if channel == "P" else "Qpf"
        bus = P_EVENT_BUS if channel == "P" else Q_EVENT_BUS
        total = float(np.sum(system.PQ.p0.v if channel == "P" else system.PQ.q0.v))
        amount = STEP_FRACTION * total
        targets = [str(idx) for idx, device_bus in zip(system.PQ.idx.v, system.PQ.bus.v)
                   if int(device_bus) == bus]
        if len(targets) != 1:
            raise RuntimeError(f"Expected one PQ device at bus {bus}; found {targets}.")
        system.add("Alter", idx=f"E3_{channel}_{sign}", t=EVENT_TIME_S,
                   model="PQ", dev=targets[0], src=source, attr="v",
                   method=sign, amount=amount)
        event = {"p_pu": amount * (1 if channel == "P" else 0),
                 "q_pu": amount * (1 if channel == "Q" else 0), "bus": float(bus)}
    system.setup()
    system.TDS.config.tf = FLAT_FINAL_TIME_S if scenario == "flat" else FINAL_TIME_S
    system.TDS.config.tstep = time_step_s
    system.TDS.config.fixt = 1
    system.TDS.config.no_tqdm = 1
    return system, event


def series_from(system: Any) -> dict[str, np.ndarray]:
    data = extract_time_series(system)
    y = system.dae.ts.y
    data["gfm_active_power_pu"] = np.asarray(y[:, system.REGF1.Pe.a], dtype=float).reshape(-1)
    data["gfm_reactive_power_pu"] = np.asarray(y[:, system.REGF1.Qe.a], dtype=float).reshape(-1)
    data["gfm_voltage_pu"] = np.asarray(y[:, system.REGF1.v.a], dtype=float).reshape(-1)
    # ``dw_y`` is an Observable (its ``a`` is not a DAE-column address).
    # Reconstruct its documented droop equation from stored state/algebraic signals.
    p_sensed = np.asarray(system.dae.ts.x[:, system.REGF1.Psen_y.a], dtype=float).reshape(-1)
    p_limited = np.asarray(y[:, system.REGF1.PIplim_y.a], dtype=float).reshape(-1)
    dw_rad_s = float(system.REGF1.w0.v[0] * system.REGF1.wdrp.v[0]) * (p_limited - p_sensed)
    dw_rad_s = np.clip(dw_rad_s, float(system.REGF1.dwmin.v[0]), float(system.REGF1.dwmax.v[0]))
    data["gfm_frequency_hz"] = NOMINAL_FREQUENCY_HZ + dw_rad_s / (2 * np.pi)
    return data


def simulate(scenario: str, time_step_s: float) -> tuple[Any, dict[str, np.ndarray], dict[str, Any]]:
    system, event = configure_system(scenario, time_step_s)
    pflow_ok = bool(system.PFlow.run())
    pflow_residual = float(np.max(np.abs(system.dae.g)))
    tds_ok = bool(system.TDS.run(no_summary=True)) if pflow_ok else False
    data = series_from(system)
    diagnostics = {
        "pflow_ok": pflow_ok, "tds_ok": tds_ok,
        "pflow_residual": pflow_residual,
        "final_residual": float(np.max(np.abs(system.dae.g))), "event": event,
    }
    return system, data, diagnostics


def metrics(data: dict[str, np.ndarray], scenario: str) -> dict[str, float]:
    t, f, v = data["time_s"], data["bus_frequency_hz"], data["bus_voltage_pu"]
    pre = t < (EVENT_TIME_S if scenario != "flat" else min(2.0, t[-1]))
    final = t >= t[-1] - 1.0
    return {
        "frequency_min_hz": float(np.min(f)), "frequency_max_hz": float(np.max(f)),
        "voltage_min_pu": float(np.min(v)), "voltage_max_pu": float(np.max(v)),
        "pre_mean_frequency_hz": float(np.mean(f[pre])),
        "final_mean_frequency_hz": float(np.mean(f[final])),
        "pre_mean_voltage_pu": float(np.mean(v[pre])),
        "final_mean_voltage_pu": float(np.mean(v[final])),
        "pre_gfm_p_pu": float(np.mean(data["gfm_active_power_pu"][pre])),
        "final_gfm_p_pu": float(np.mean(data["gfm_active_power_pu"][final])),
        "pre_gfm_q_pu": float(np.mean(data["gfm_reactive_power_pu"][pre])),
        "final_gfm_q_pu": float(np.mean(data["gfm_reactive_power_pu"][final])),
        "flat_frequency_drift_hz": float(np.max(np.abs(f[-1] - f[0]))),
        "flat_voltage_drift_pu": float(np.max(np.abs(v[-1] - v[0]))),
    }


def run_eig() -> dict[str, Any]:
    system, _ = configure_system("flat", TIME_STEP_S)
    pflow_ok = bool(system.PFlow.run())
    eig_ok = bool(system.EIG.run()) if pflow_ok else False
    eigenvalues = np.asarray(system.EIG.mu, dtype=complex)
    finite = eigenvalues[np.isfinite(eigenvalues)]
    return {
        "pflow_ok": pflow_ok, "eig_ok": eig_ok, "eigenvalue_count": int(finite.size),
        "maximum_real_part": float(np.max(finite.real)),
        "unstable_mode_count": int(np.sum(finite.real > 1e-6)),
        "near_zero_mode_count": int(np.sum(np.abs(finite.real) <= 1e-8)),
        "eigenvalues_real": finite.real.tolist(), "eigenvalues_imag": finite.imag.tolist(),
    }


def plot_frequency(all_data: dict[str, dict[str, np.ndarray]]) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(11, 7), sharex=True, sharey=True)
    for axis, (name, data) in zip(axes.flat, all_data.items()):
        ids, f = data["frequency_bus_ids"], data["bus_frequency_hz"]
        critical = np.isin(ids, CRITICAL_BUSES)
        axis.fill_between(data["time_s"], f.min(1), f.max(1), color="0.85", label="39-bus envelope")
        for col in np.flatnonzero(critical):
            axis.plot(data["time_s"], f[:, col], lw=.8, label=f"Bus {ids[col]}")
        if "gfm_frequency_hz" in data:
            axis.plot(data["time_s"], data["gfm_frequency_hz"], "k--", lw=1.2, label="GFM32")
        axis.axhline(60, color="black", ls=":", lw=.8)
        axis.axvline(EVENT_TIME_S, color="tab:red", ls=":", lw=.8)
        axis.set_title(name); axis.grid(alpha=.25); axis.ticklabel_format(axis="y", style="plain", useOffset=False)
    axes[0, 0].legend(ncol=3, fontsize=6)
    for axis in axes[:, 0]: axis.set_ylabel("Frequency [Hz]")
    for axis in axes[-1, :]: axis.set_xlabel("Time [s]")
    fig.suptitle("E3--C0a: filtered bus frequency and GFM32 frequency")
    fig.tight_layout(); fig.savefig(FREQUENCY_FIGURE, dpi=300, bbox_inches="tight"); plt.close(fig)


def plot_voltage(all_data: dict[str, dict[str, np.ndarray]]) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(11, 7), sharex=True)
    for axis, (name, data) in zip(axes.flat, all_data.items()):
        ids, v = data["bus_ids"], data["bus_voltage_pu"]
        axis.fill_between(data["time_s"], v.min(1), v.max(1), color="0.85", label="39-bus envelope")
        for col in np.flatnonzero(np.isin(ids, CRITICAL_BUSES)):
            axis.plot(data["time_s"], v[:, col], lw=.8, label=f"Bus {ids[col]}")
        axis.axvline(EVENT_TIME_S, color="tab:red", ls=":", lw=.8)
        axis.set_title(name); axis.grid(alpha=.25)
    axes[0, 0].legend(ncol=4, fontsize=6)
    for axis in axes[:, 0]: axis.set_ylabel("Voltage |V| [pu]")
    for axis in axes[-1, :]: axis.set_xlabel("Time [s]")
    fig.suptitle("E3--C0a: bus-voltage response")
    fig.tight_layout(); fig.savefig(VOLTAGE_FIGURE, dpi=300, bbox_inches="tight"); plt.close(fig)


def plot_power(all_data: dict[str, dict[str, np.ndarray]]) -> None:
    fig, axes = plt.subplots(2, 1, figsize=(9, 7), sharex=True)
    for name, data in all_data.items():
        axes[0].plot(data["time_s"], data["gfm_active_power_pu"], label=name)
        axes[1].plot(data["time_s"], data["gfm_reactive_power_pu"], label=name)
    axes[0].set_ylabel("GFM32 P [pu]"); axes[1].set_ylabel("GFM32 Q [pu]")
    axes[1].set_xlabel("Time [s]")
    for axis in axes:
        axis.axvline(EVENT_TIME_S, color="tab:red", ls=":"); axis.grid(alpha=.25); axis.legend()
    fig.suptitle("E3--C0a: delivered GFM32 power")
    fig.tight_layout(); fig.savefig(POWER_FIGURE, dpi=300, bbox_inches="tight"); plt.close(fig)


def update_matrix(summary: dict[str, Any]) -> None:
    with MATRIX_FILE.open("r", newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream); fields = list(reader.fieldnames or [])
        rows = [row for row in reader if row.get("stage") != "E3"]
    row = {key: "" for key in fields}
    row.update({
        "run_id": summary["run_id"], "stage": "E3", "case": "C0a_9SG_1GFM",
        "scenario": "flat_and_small_signed_P_Q_steps", "baseline": "E2_C0",
        "source_sha256": summary["source_sha256"], "variant_sha256": summary["variant_sha256"],
        "andes_version": summary["andes_version"], "git_commit": summary["git_commit"],
        "seed": "not_applicable_deterministic", "simulation_t0": 0.0, "simulation_tf": FINAL_TIME_S,
        "solver_settings": json.dumps(summary["solver"], sort_keys=True), "event_time": EVENT_TIME_S,
        "event_bus": "4(P);21(Q)", "event_P_MW": summary["events"]["P+_mw"],
        "event_Q_MVAr": summary["events"]["Q+_mvar"], "fn_Hz": NOMINAL_FREQUENCY_HZ,
        "bus_freq_set": "1-39", "bus_volt_set": "1-39",
        "f_nadir_Hz": summary["metrics"]["P+"]["frequency_min_hz"],
        "min_V_pu": min(v["voltage_min_pu"] for v in summary["metrics"].values()),
        "max_V_pu": max(v["voltage_max_pu"] for v in summary["metrics"].values()),
        "dae_residual": max(v["final_residual"] for v in summary["diagnostics"].values()),
        "status_auto": summary["status_auto"], "status_human": "PENDING",
        "frequency_figure": FREQUENCY_FIGURE.relative_to(PROJECT_ROOT).as_posix(),
        "voltage_figure": VOLTAGE_FIGURE.relative_to(PROJECT_ROOT).as_posix(),
        "raw_file": RAW_FILE.relative_to(PROJECT_ROOT).as_posix(),
        "decision_note": "E4 remains blocked pending human approval of E3 figures.",
    })
    with MATRIX_FILE.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields); writer.writeheader(); writer.writerows(rows + [row])


def main() -> int:
    build_variant()
    scenario_defs = {"P+": ("P", "+"), "P-": ("P", "-"),
                     "Q+": ("Q", "+"), "Q-": ("Q", "-")}
    flat_system, flat_data, flat_diag = simulate("flat", TIME_STEP_S)
    all_data: dict[str, dict[str, np.ndarray]] = {}
    diagnostics: dict[str, Any] = {"flat": flat_diag}
    metric_values: dict[str, Any] = {"flat": metrics(flat_data, "flat")}
    for name, definition in scenario_defs.items():
        _, data, diag = simulate(definition, TIME_STEP_S)
        all_data[name] = data; diagnostics[name] = diag; metric_values[name] = metrics(data, name)
    refined: dict[str, Any] = {}
    for name in ("P+", "Q+"):
        _, data, diag = simulate(scenario_defs[name], REFINED_TIME_STEP_S)
        refined[name] = {"diagnostics": diag, "metrics": metrics(data, name)}
    eig = run_eig()
    flat_ok = metric_values["flat"]["flat_frequency_drift_hz"] <= 1e-4 and metric_values["flat"]["flat_voltage_drift_pu"] <= 1e-5
    signs_ok = (metric_values["P+"]["frequency_min_hz"] < metric_values["P-"]["frequency_min_hz"]
                and metric_values["Q+"]["final_mean_voltage_pu"] < metric_values["Q-"]["final_mean_voltage_pu"])
    numerical = {
        name: {"frequency_min_abs_difference_hz": abs(metric_values[name]["frequency_min_hz"] - refined[name]["metrics"]["frequency_min_hz"]),
               "voltage_min_abs_difference_pu": abs(metric_values[name]["voltage_min_pu"] - refined[name]["metrics"]["voltage_min_pu"])}
        for name in refined
    }
    numerical_ok = all(v["frequency_min_abs_difference_hz"] <= 0.002 and v["voltage_min_abs_difference_pu"] <= 0.002 for v in numerical.values())
    all_runs_ok = all(d["pflow_ok"] and d["tds_ok"] and d["final_residual"] <= 1e-5 for d in diagnostics.values())
    structure_ok = len(flat_system.GENROU.idx.v) == 9 and len(flat_system.REGF1.idx.v) == 1 and "GENROU_3" not in flat_system.GENROU.idx.v
    automatic_pass = all((flat_ok, signs_ok, numerical_ok, all_runs_ok, structure_ok,
                          eig["eig_ok"], eig["unstable_mode_count"] == 0))
    status = "PASS_AUTO_WAIT_HUMAN" if automatic_pass else "FAIL"
    raw: dict[str, np.ndarray] = {}
    for name, data in {"flat": flat_data, **all_data}.items():
        for key, value in data.items(): raw[f"{name}_{key}"] = value
    np.savez_compressed(RAW_FILE, **raw)
    plot_frequency(all_data); plot_voltage(all_data); plot_power(all_data)
    summary = {
        "run_id": f"IEEE39_E3_C0a_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}",
        "created_utc": datetime.now(timezone.utc).isoformat(), "status_auto": status,
        "status_human": "PENDING", "source_case": str(E0_CASE.relative_to(PROJECT_ROOT)),
        "variant_case": str(VARIANT_CASE.relative_to(PROJECT_ROOT)),
        "source_sha256": sha256(E0_CASE), "variant_sha256": sha256(VARIANT_CASE),
        "runner_sha256": sha256(Path(__file__).resolve()), "git_commit": git_commit(),
        "python_version": platform.python_version(), "andes_version": andes.__version__,
        "configuration": {"sg_count": 9, "gfm_count": 1, "gfm_model": "REGF1",
                          "gfm_bus": GFM_BUS, "static_gen": GFM_STATIC_GEN,
                          "Sn_MVA": GFM_SN_MVA, "removed_dynamic_rows": REMOVED_DYNAMIC_ROWS,
                          "parameter_policy": "ANDES 2.0.0 REGF1 defaults except bus, gen, Sn, fn"},
        "events": {"P+_mw": diagnostics["P+"]["event"]["p_pu"] * flat_system.config.mva,
                   "P-_mw": -diagnostics["P-"]["event"]["p_pu"] * flat_system.config.mva,
                   "Q+_mvar": diagnostics["Q+"]["event"]["q_pu"] * flat_system.config.mva,
                   "Q-_mvar": -diagnostics["Q-"]["event"]["q_pu"] * flat_system.config.mva},
        "solver": {"method": flat_system.TDS.config.method, "tf_flat_s": FLAT_FINAL_TIME_S,
                   "tf_step_s": FINAL_TIME_S, "time_step_s": TIME_STEP_S,
                   "refined_time_step_s": REFINED_TIME_STEP_S,
                   "tolerance": flat_system.TDS.config.tol,
                   "maximum_iterations": flat_system.TDS.config.max_iter},
        "diagnostics": diagnostics, "metrics": metric_values, "refined": refined,
        "numerical_step_comparison": numerical, "eigenanalysis": eig,
        "checks": {"flat_no_drift": flat_ok, "response_signs": signs_ok,
                   "numerical_convergence": numerical_ok, "all_runs_converged": all_runs_ok,
                   "structure_9sg_1gfm_no_duplicate": structure_ok,
                   "no_positive_eigenvalues": eig["unstable_mode_count"] == 0},
        "limitations": ["REGF1 is a positive-sequence RMS model and contains no BESS state-of-charge model.",
                        "EIG found no positive-real-part modes, but exact/near-zero modes and an ill-conditioning warning prevent a strict Hurwitz claim.",
                        "The signed 1% load steps are diagnostic and do not establish nonlinear safety."],
    }
    SUMMARY_FILE.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    update_matrix(summary)
    print(json.dumps({"status_auto": status, "checks": summary["checks"],
                      "events": summary["events"], "eigenanalysis": {k: v for k, v in eig.items() if not k.startswith("eigenvalues_")}}, indent=2))
    return 0 if automatic_pass else 1


if __name__ == "__main__":
    raise SystemExit(main())
