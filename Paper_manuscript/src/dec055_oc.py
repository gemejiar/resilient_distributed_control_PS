"""DEC-055 operating conditions OC1/OC2/OC3: case generation (new files only; the frozen C2t case is never written), case selection inside the
existing runners (module attribute patch, no edit of historical code), and an OC-aware plant / pooled plant for the static stages of L0.

    <venv python> dec055_oc.py build      -> writes the OC2 and OC3 case files exclusively and records their hashes"""
from __future__ import annotations

import json
import os
import sys
from concurrent.futures import ProcessPoolExecutor
from concurrent.futures.process import BrokenProcessPool

import numpy as np

import dec054_common as C54
import dec055_common as C

ENV = "DEC055_OC"


def scale_rule(p0_pv, pmax_pv, s):
    """PV p0 of the OC reference point: s * p0, capped at pmax - PV_CAP_BACKOFF_PU (declared rule)."""
    return np.minimum(np.asarray(p0_pv, float) * s, np.asarray(pmax_pv, float) - C.PV_CAP_BACKOFF_PU)


def build_case(oc: str, src=None, dst=None) -> dict:
    """Copy the frozen case workbook and scale PQ p0/q0 (power factor preserved) and PV p0 (rule above).  Exclusive creation."""
    import openpyxl
    src = src or C.BASE_CASE
    dst = dst or C.OC[oc]["case"]
    s = C.OC[oc]["load_scale"]
    h0 = C54.sha_file(src)
    wb = openpyxl.load_workbook(src)
    changed = {"PQ": [], "PV": []}
    for sheet, cols in (("PQ", ("p0", "q0")), ("PV", ("p0",))):
        ws = wb[sheet]
        head = [c.value for c in ws[1]]
        ix = {k: head.index(k) for k in set(cols) | {"idx", "bus", "pmax"} if k in head}
        for row in ws.iter_rows(min_row=2):
            if row[ix["idx"]].value is None:
                continue
            if sheet == "PQ":
                for k in cols:
                    old = float(row[ix[k]].value)
                    row[ix[k]].value = old * s
                    changed["PQ"].append({"bus": int(row[ix["bus"]].value), "col": k, "old": old, "new": old * s})
            else:
                old = float(row[ix["p0"]].value)
                new = float(scale_rule([old], [float(row[ix["pmax"]].value)], s)[0])
                row[ix["p0"]].value = new
                changed["PV"].append({"bus": int(row[ix["bus"]].value), "old": old, "new": new, "capped": bool(new < old * s - 1e-12)})
    with open(dst, "xb") as f:
        wb.save(f)
    assert C54.sha_file(src) == h0, "the source case changed"
    return {"oc": oc, "load_scale": s, "source": _rel(src), "source_sha256": h0, "case": _rel(dst), "case_sha256": C54.sha_file(dst), "changes": changed}


def _rel(p):
    try:
        return str(p.relative_to(C.ROOT))
    except ValueError:
        return str(p)


def select(oc: str):
    """Point the existing runners at the OC case (module attributes only).  Must be called before any plant is built in this process."""
    os.environ.setdefault("C2_VARIANT", "C2t")
    os.environ[ENV] = oc
    import run_e5_ieee39 as e5
    import run_e8_ieee39 as e8
    import select_equilibrium_b3 as b3
    case = C.OC[oc]["case"]
    if not case.exists():
        raise FileNotFoundError(case)
    e5.CASE = e8.CASE = b3.CASE = case
    if not getattr(e8.configure, "_dec055", False):           # D1 (declared before freezing): PFlow tolerance 1e-10 for every OC (ANDES default 1e-6 left
        orig = e8.configure                                   # |g| = 1.13e-8 after TDS.init at the OC2 reference point, above the 1e-8 validity threshold)

        def configure(scenario, time_step_s, _orig=orig):
            system, ctx = _orig(scenario, time_step_s)
            system.PFlow.config.tol = PFLOW_TOL
            return system, ctx
        configure._dec055 = True
        e8.configure = configure
    return case


PFLOW_TOL = 1e-10


def current_oc() -> str:
    return os.environ.get(ENV, "OC1")


def make_plant(oc: str):
    import dec054_plant as PL
    select(oc)
    p = PL.AndesPlant(require_backup=False)        # L0 is static (no TDS): the TDS ledger/gate is not involved
    p.oc = oc
    return p


# ---------------------------------------------------------------- OC-aware pool (same interface as dec054_pool.PooledPlant)
def _init_worker(oc):
    import dec054_pool as POOL
    os.environ.update({v: "1" for v in ("OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "OMP_NUM_THREADS", "NUMEXPR_NUM_THREADS")})
    POOL._W = make_plant(oc)


class OCPooledPlant:
    """dec054_pool.PooledPlant with the OC case in the local plant and in every worker (also after a pool rebuild)."""

    def __init__(self, oc: str, workers: int = 3):
        import dec054_pool as POOL
        for v in ("OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "OMP_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
            os.environ[v] = "1"
        self.POOL, self.oc, self.workers = POOL, oc, workers
        self.local = make_plant(oc)
        for a in ("limits", "z_ref", "z_lo", "z_hi", "var_names"):
            setattr(self, a, getattr(self.local, a))
        self.units = self.local.units
        self._pool = ProcessPoolExecutor(max_workers=workers, initializer=_init_worker, initargs=(oc,))
        self._ev, self._evt = {}, {}
        self.n_event_tasks = self.pool_rebuilds = self.n_eval = 0

    close = lambda self: self._pool.shutdown(wait=True, cancel_futures=True)          # noqa: E731

    def evaluate(self, z):
        k = self.POOL._key(z)
        if k not in self._ev:
            self.n_eval += 1
            self._ev[k] = self.local.evaluate(np.asarray(z, float))
        return self._ev[k]

    def prefetch(self, zs, scs, exc: bool = True):
        P = self.POOL
        todo, keys = [], []
        for z in zs:
            if not self.evaluate(z)["ok"]:
                continue
            for sc in scs:
                k = (P._key(z), P._skey(sc))
                have = self._evt.get(k)
                if (have is None or (exc and have["exc"] is None and have["jump"].get("ok"))) and k not in keys:
                    keys.append(k)
                    todo.append((np.asarray(z, float).tolist(), P._skey(sc), exc))
        if todo:
            self.n_event_tasks += len(todo)
            for attempt in (0, 1):
                try:
                    for k, r in zip(keys, self._pool.map(P._task, todo, chunksize=1)):
                        self._evt[k] = r
                    break
                except BrokenProcessPool:
                    self.pool_rebuilds += 1
                    self._pool.shutdown(wait=False, cancel_futures=True)
                    self._pool = ProcessPoolExecutor(max_workers=self.workers, initializer=_init_worker, initargs=(self.oc,))
                    if attempt == 1:
                        for k in keys:
                            self._evt.setdefault(k, {"jump": {"ok": False, "reason": "BrokenProcessPool"}, "exc": {"ok": False}})

    def _event(self, z, sc, exc):
        k = (self.POOL._key(z), self.POOL._skey(sc))
        have = self._evt.get(k)
        if have is None or (exc and have["exc"] is None and have["jump"].get("ok")):
            self.prefetch([z], [sc], exc)
        return self._evt.get(k, {"jump": {"ok": False}, "exc": {"ok": False}})

    def jump(self, z, sc):
        return self._event(z, sc, False)["jump"]

    def excursion(self, z, sc):
        r = self._event(z, sc, True)["exc"]
        return r if r is not None else {"ok": False}


if __name__ == "__main__" and sys.argv[1:] == ["build"]:
    out = [build_case(oc) for oc in ("OC2", "OC3")]
    with open(C.ART / "IEEE39_DEC055_oc_cases.json", "x", encoding="utf-8") as f:
        json.dump({"stage": "DEC-055 L0 OC case generation (no TDS)", "rule": "PQ p0,q0 x s; PV p0 = min(s p0, pmax - %.2f); v0, Sn, Vn, dynamics, topology unchanged" % C.PV_CAP_BACKOFF_PU,
                   "OC1": {"case": str(C.BASE_CASE.relative_to(C.ROOT)), "case_sha256": C54.sha_file(C.BASE_CASE), "load_scale": 1.0}, "OC2": out[0], "OC3": out[1]}, f, indent=1)
    print(json.dumps([{k: o[k] for k in ("oc", "case_sha256")} for o in out]))
